"""Conservative extraction, with optional explicitly enabled OCR/LLM."""

import json
import re
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path

from ..schema import Event, Evidence
from ..security import image_uri, local_file

LABELS = {
    "title": ["活动名称", "名称", "标题", "title"],
    "category": ["活动类型", "类别", "category"],
    "start": ["开始时间", "开始", "start"],
    "end": ["结束时间", "结束", "end"],
    "timezone": ["时区", "timezone"],
    "venue": ["地点", "地址", "venue"],
    "online_platform": ["线上平台", "平台", "online_platform"],
    "organizer": ["主办方", "organizer"],
    "audience": ["受众", "audience"],
    "highlights": ["亮点", "highlights"],
    "speakers": ["主讲人", "speakers"],
    "fee": ["费用", "fee"],
    "registration": ["报名方式", "报名", "registration"],
    "registration_url": ["报名链接", "registration_url"],
    "deadline": ["截止时间", "deadline"],
    "contact": ["联系方式", "contact"],
    "subtitle": ["副标题", "subtitle"],
}

CATEGORIES = {
    "wedding": ["婚礼", "wedding"], "dinner": ["晚宴", "gala", "dinner"],
    "awards": ["颁奖", "awards"], "sports": ["足球", "篮球", "运动会", "跑步", "sports"],
    "esports": ["电竞", "esports"], "night-party": ["DJ", "音乐派对", "夜间派对"],
    "academic": ["学术", "academic"], "workshop": ["研讨", "工坊", "workshop"],
    "recruitment": ["招新", "recruitment"], "market": ["市集", "market"],
    "exhibition": ["展览", "exhibition"], "launch": ["发布", "launch"],
    "opening": ["开业", "opening"], "birthday": ["生日", "birthday"],
    "forum": ["论坛", "forum"], "lecture": ["讲座", "分享", "lecture"],
    "community": ["社群", "community"],
}


def checked_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        if sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024:
            raise ValueError("DOCX expanded size exceeds 80 MiB")


def read_text(path: Path, ocr: bool = False) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".docx":
        from docx import Document
        checked_zip(path)
        document = Document(path)
        rows = [p.text for p in document.paragraphs]
        rows += ["\n".join(": ".join(c.text for c in r.cells) for r in t.rows) for t in document.tables]
        return "\n".join(rows)
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if len(reader.pages) > 100:
            raise ValueError("PDF exceeds 100 pages")
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        if not text.strip():
            raise ValueError("Scanned PDF has no text layer; use the Agent's vision or transcribe it")
        if len(text) > 400_000:
            raise ValueError("PDF text exceeds the extraction limit")
        # Browser PDF font mappings may use compatibility radicals/ligatures.
        # Normalize the extracted text rather than losing equivalent event facts.
        return unicodedata.normalize("NFKC", text)
    if suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        image_uri(path)
        if not ocr:
            raise ValueError("Screenshot extraction requires --ocr and Tesseract, or Agent vision")
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path), lang="chi_sim+eng", timeout=30)
    raise ValueError("Supported inputs: JSON, UTF-8 TXT/Markdown, DOCX, PDF, PNG/JPEG/WebP")


def parse_text(text: str, source: str = "text") -> Event:
    data: dict = {"sources": {}, "uncertainties": []}
    for field, aliases in LABELS.items():
        matches = re.findall(r"(?im)^\s*(?:[-*#]\s*)?(?:" + "|".join(aliases) + r")\s*[:：]\s*(.+)$", text)
        if not matches:
            continue
        distinct = list(dict.fromkeys(m.strip() for m in matches))
        if len(distinct) > 1:
            data["uncertainties"].append(f"{field} 存在冲突，请确认 / conflicting values")
        value = distinct[0]
        if field in {"audience", "highlights", "speakers"}:
            value = [v.strip() for v in re.split("[、;；|]", value) if v.strip()]
        if field in {"start", "end", "deadline"}:
            try:
                datetime.fromisoformat(value)
                if "T" not in value and " " not in value:
                    raise ValueError()
            except ValueError:
                data["uncertainties"].append(f"{field} 需要 ISO 日期时间 / ISO date and time required")
                continue
        data[field] = value
        data["sources"][field] = Evidence(source=source, excerpt=str(value), confidence=0.95, status="extracted")
    if not data.get("title"):
        heading = re.search(r"(?m)^#\s+(.+)$", text)
        if heading:
            data["title"] = heading[1]
            data["sources"]["title"] = Evidence(source=source, excerpt=heading[1], confidence=0.85, status="extracted")
    if not data.get("start"):
        date = re.search(r"(20\d{2})[年/.-](\d{1,2})[月/.-](\d{1,2})(?:日)?[^\n]{0,12}?(\d{1,2})[:：点](\d{2})?", text)
        if date:
            try:
                stamp = datetime(*map(int, date.groups()[:4]), int(date[5] or 0)).isoformat(timespec="minutes")
                data["start"] = stamp
                data["sources"]["start"] = Evidence(source=source, excerpt=date[0], confidence=0.8, status="extracted")
            except ValueError:
                data["uncertainties"].append("原文日期无效 / Invalid source date")
        if re.search("下个月|明天|后天|下周|next month|tomorrow", text, re.I):
            data["uncertainties"].append("相对日期需要确认准确年月日 / Confirm the relative date")
    if not data.get("registration_url"):
        urls = re.findall(r"https?://[^\s<>）)]+", text)
        if len(urls) == 1 and re.search(r"报名|注册|RSVP|register", text, re.I):
            data["registration_url"] = urls[0].rstrip("，。")
            data["sources"]["registration_url"] = Evidence(source=source, excerpt=urls[0], confidence=0.7, status="extracted")
    if not data.get("fee") and "免费" in text:
        data["fee"] = "免费"
        data["sources"]["fee"] = Evidence(source=source, excerpt="免费", confidence=0.95, status="extracted")
    if "category" not in data:
        for category, words in CATEGORIES.items():
            found = next((word for word in words if word.lower() in text.lower()), None)
            if found:
                data["category"] = category
                data["sources"]["category"] = Evidence(source=source, excerpt=found, confidence=0.75, status="inferred")
                break
    # Preferences may be inferred; audience, organizer, venue and fees may not.
    if data.get("category") in {"wedding", "dinner", "awards"}:
        data["formality"] = "ceremonial"
        data["sources"]["formality"] = Evidence(source=source, excerpt=data["category"], confidence=0.7, status="inferred")
    elif data.get("category") in {"academic", "forum", "lecture", "workshop"}:
        data["formality"] = "professional"
        data["sources"]["formality"] = Evidence(source=source, excerpt=data["category"], confidence=0.7, status="inferred")
    data["keywords"] = [word for word in ["AI", "Agent"]
                        if re.search(r'(?<![A-Za-z0-9])'+word+r'(?![A-Za-z0-9])',text,re.I)]
    data['keywords'] += [word for word in ['科技','艺术','手作','音乐','设计'] if word in text]
    data["demo"] = bool(re.search(r"虚构活动|fictional event|^demo\s*[:：]\s*true", text, re.I | re.M))
    return Event.model_validate(data)


def extract(path: str | Path, root: Path | None = None, ocr: bool = False) -> Event:
    path = local_file(path, root)
    if path.suffix.lower() == ".json":
        raw = json.loads(path.read_text(encoding="utf-8-sig"))
        event = Event.model_validate(raw)
        if event.assets.visual_receipt:
            asset_path = Path(event.assets.visual_receipt)
            if not asset_path.is_absolute():
                event.assets.visual_receipt = str((path.parent / asset_path).resolve())
        for ratio, value in event.assets.visual_variants.items():
            asset_path = Path(value)
            if not asset_path.is_absolute():
                event.assets.visual_variants[ratio] = str((path.parent / asset_path).resolve())
        for key, value in raw.items():
            if value and key not in {"sources", "confirmed_fields", "uncertainties", "schema_version"}:
                event.sources.setdefault(key, Evidence(source=path.name, confidence=1))
        return event
    return parse_text(read_text(path, ocr), source=path.name)
